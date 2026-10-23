"""Independent openpyxl + scipy-exact audit of the frozen HAP1 pilot."""
from pathlib import Path
import json
import sys
import numpy as np
import pandas as pd
import openpyxl
from scipy.stats import permutation_test

out=Path(sys.argv[1]); spec=json.loads((out/'R15_HAP1_FROZEN_SPEC.json').read_text())
full=pd.read_csv(out/'S4_ARID1A_all_gene_scores.csv').set_index('gene')
fixed=sorted(set(g for members in spec['fixed_sets'].values() for g in members))
book=openpyxl.load_workbook(out/'raw/File_S4.xlsx',read_only=True,data_only=True)
checks=[]
for sheet,colname in [('qGI_scores','author_qGI'),('qGI_FDR','author_FDR')]:
    print('INDEPENDENT_OPENPYXL',sheet,flush=True)
    ws=book[sheet]; rows=ws.iter_rows(values_only=True); header=next(rows)
    i=header.index('ARID1A_019_min'); gi=header.index('Gene.symbol')
    remaining=set(fixed+['A1BG','A1CF','A2ML1'])
    for r in rows:
        gene=spec['aliases'].get(r[gi],r[gi])
        if gene not in remaining: continue
        expected=float(full.loc[gene,colname]); observed=float(r[i])
        checks.append({'sheet':sheet,'gene':gene,'original_excel':observed,'extracted_csv':expected,
                       'absolute_error':abs(observed-expected)})
        remaining.remove(gene)
        if not remaining: break
    assert not remaining,(sheet,remaining)
check=pd.DataFrame(checks); check.to_csv(out/'independent_excel_value_checks.csv',index=False)
assert check.absolute_error.max()<1e-12

# The author pair table contains listed significant interactions, not all
# 5.3 million screened pairs; do not infer a score of zero from an absent row.
print('INDEPENDENT_AUTHOR_PAIR_TABLE',flush=True)
ws=book['pairwise GI_Complete dataset']; rows=ws.iter_rows(values_only=True); header=next(rows)
listed=[]
for r in rows:
    obj=dict(zip(header,r))
    if obj['query_screen']!='ARID1A_019_min': continue
    gene=spec['aliases'].get(obj['library_gene'],obj['library_gene']); obj['gene']=gene
    obj['qGI_matrix_abs_error']=abs(float(obj['qGI_score'])-full.loc[gene,'author_qGI'])
    obj['FDR_matrix_abs_error']=abs(float(obj['FDR'])-full.loc[gene,'author_FDR'])
    listed.append(obj)
listed=pd.DataFrame(listed); listed.to_csv(out/'independent_author_listed_pairs_checks.csv',index=False)
assert listed.qGI_matrix_abs_error.max()<1e-12
assert listed.FDR_matrix_abs_error.max()<1e-12
book.close()

c=pd.read_csv(out/'primary_direct_contrasts.csv'); permutation_checks=[]
for row in c.to_dict('records'):
    a=-full.loc[spec['fixed_sets'][row['group_A']],'author_qGI'].to_numpy()
    b=-full.loc[spec['fixed_sets'][row['group_B']],'author_qGI'].to_numpy()
    result=permutation_test((a,b),lambda x,y:np.mean(x)-np.mean(y),n_resamples=np.inf,
                            vectorized=False,permutation_type='independent',alternative='greater')
    # Match the preregistered absolute-statistic two-sided definition, which can
    # differ from SciPy's default doubled smaller tail for unequal group sizes.
    p=float(np.mean(abs(result.null_distribution)>=abs(result.statistic)-1e-12))
    permutation_checks.append({'group_A':row['group_A'],'group_B':row['group_B'],
      'scipy_observed_difference':float(result.statistic),'scipy_null_count':len(result.null_distribution),
      'independent_abs_two_sided_p':p,'primary_p':row['p_two_sided_exact'],
      'absolute_p_error':abs(p-row['p_two_sided_exact'])})
pc=pd.DataFrame(permutation_checks); pc.to_csv(out/'independent_scipy_permutation_checks.csv',index=False)
assert pc.absolute_p_error.max()<1e-12
p=c.p_two_sided_exact.to_numpy(); order=np.argsort(p)
holm=np.empty(len(p)); holm[order]=np.minimum(1,np.maximum.accumulate(p[order]*np.arange(len(p),0,-1)))
assert np.max(abs(holm-c.Holm_p_three_prespecified.to_numpy()))<1e-12

summary={'status':'PASS','excel_value_checks':len(check),'max_excel_absolute_error':float(check.absolute_error.max()),
 'author_ARID1A_listed_pairs':len(listed),'author_pair_table_qGI_max_error':float(listed.qGI_matrix_abs_error.max()),
 'author_pair_table_FDR_max_error':float(listed.FDR_matrix_abs_error.max()),
 'direct_contrasts_independently_recomputed':len(pc),'max_exact_permutation_p_error':float(pc.absolute_p_error.max()),
 'holm_independent_max_error':float(np.max(abs(holm-c.Holm_p_three_prespecified.to_numpy()))),
 'FDR_warning':'qGI_FDR is author-provided BH FDR; S4 does not include gene raw P values, so original genome-wide BH cannot be independently regenerated from this workbook. Exact stored FDR values and three-comparison Holm were independently checked.',
 'pair_table_warning':'The sheet labelled pairwise GI_Complete dataset contains 131680 listed interaction rows, not the complete 17804 by 298 screen-by-library-gene matrix (5305592 measurements, including repeated query screens). Full unfiltered matrices were used for our denominator and tests.'}
(out/'R15_HAP1_INDEPENDENT_AUDIT.json').write_text(json.dumps(summary,indent=2),encoding='utf-8')
print(json.dumps(summary,indent=2),flush=True)
