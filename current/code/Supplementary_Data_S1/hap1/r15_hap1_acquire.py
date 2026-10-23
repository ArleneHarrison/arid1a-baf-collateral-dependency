"""Acquire only the public HAP1 metadata and fixed-panel GI inputs on seeta.

Stage 1 freezes member lists and analyses before reading effect columns.
No scientific computation is run on the local workstation.
"""
from __future__ import annotations
import argparse
import hashlib
import json
from pathlib import Path
import platform
import time
import urllib.request
import openpyxl

BASE = 'https://boonelab.ccbr.utoronto.ca/supplement/billmanncostanzo2026/'
SETS = {
    'contact6': ['SMARCC2','SMARCA4','SMARCD1','SMARCB1','SMARCE1','DPF2'],
    'broad_BAF_background6': ['ACTL6A','PBRM1','ARID2','BRD7','PHF10','BRD9'],
    'cBAF_associated9': ['ARID1B','DPF1','DPF2','DPF3','BCL7A','BCL7B','BCL7C','SS18','SS18L1'],
    'PBAF4': ['PBRM1','ARID2','BRD7','PHF10'],
    'ncBAF3': ['BRD9','BICRA','BICRAL'],
}

def sha(path):
    h = hashlib.sha256()
    with Path(path).open('rb') as f:
        for b in iter(lambda: f.read(2**20), b''): h.update(b)
    return h.hexdigest()

def save(path, obj):
    Path(path).write_text(json.dumps(obj, indent=2, ensure_ascii=False, default=str), encoding='utf-8')

def fetch(url, target, limit=600_000_000):
    req = urllib.request.Request(url, headers={'User-Agent':'ARID1A-BAF-scientific-reanalysis/1.0'})
    started = time.time()
    with urllib.request.urlopen(req, timeout=120) as res:
        headers = dict(res.headers)
        size = int(res.headers.get('Content-Length',0))
        if size > limit: raise ValueError(f'File exceeds bounded download limit: {size}')
        count = 0
        with target.open('wb') as out:
            while True:
                block = res.read(2**20)
                if not block: break
                out.write(block); count += len(block)
                if count > limit: raise ValueError('Stream exceeds limit')
    return {'url':url,'path':str(target),'bytes':target.stat().st_size,'sha256':sha(target),
            'http_headers':headers,'download_seconds':time.time()-started,
            'acquired_utc':time.strftime('%Y-%m-%dT%H:%M:%SZ',time.gmtime())}

def main():
    ap=argparse.ArgumentParser(); ap.add_argument('--out',required=True); args=ap.parse_args()
    out=Path(args.out); raw=out/'raw'; raw.mkdir(parents=True,exist_ok=True)
    specfile=out/'R15_HAP1_FROZEN_SPEC.json'
    spec={
      'stage':'frozen_before_reading_new_effect_columns',
      'frozen_utc':time.strftime('%Y-%m-%dT%H:%M:%SZ',time.gmtime()),
      'anchor':'ARID1A','primary_query':'GIN019','fixed_sets':SETS,
      'source_of_sets':'R14 manuscript Methods: 6-contact and 6 broad-BAF protein panel; nine-gene cBAF-associated operational module; R10/R11 frozen PBAF/ncBAF modules.',
      'aliases':{'GLTSCR1':'BICRA','GLTSCR1L':'BICRAL'},
      'primary_effect':'author qGI from S4, negative means stronger loss-conditioned depletion; analyse -qGI so positive means greater dependency',
      'primary_contrasts':[['contact6','broad_BAF_background6'],['cBAF_associated9','PBAF4'],['cBAF_associated9','ncBAF3']],
      'primary_statistic':'difference of mean -qGI; exhaustive two-sided gene-label permutation; Holm correction of 3 prespecified comparisons',
      'coverage_rule':'Primary contrasts require full fixed-set coverage with finite author qGI and author FDR. Partial coverage is descriptive only; never replace members.',
      'gene_significance':'retain author gene-level FDR; standard author hit threshold qGI<-0.3 and FDR<0.1; do not recompute FDR only among BAF genes',
      'optional_enrichment':'Only full-coverage sets; one-sided Fisher exact for negative author hits against all actual finite scored ARID1A targets excluding ARID1A, not genome annotation; BH over 5 frozen sets; secondary competitive test.',
      'replicate_unit':'One HAP1 ARID1A clone/screen, three technical endpoints are not biological n. Permuted unit is gene; gene exchangeability is a competitive gene-set assumption and does not establish biological replication.',
      'validation_S6':'Alternative guides in 1215 hit-selected targets; not unbiased independently selected genome-wide validation. Preserve all available cases and missingness. Compare author validation dLFC to primary -qGI only descriptively when fixed-set coverage is incomplete; no module claim or invented FDR.',
      'GO_rule':'Data-complete reproducible GI extraction plus at least one positive full-panel primary contrast surviving 3-test Holm can support bounded supplemental strengthening only; no automatic manuscript upgrade. Stronger independent cBAF specificity requires both direct contrasts and independent biological models; one HAP1 clone cannot establish NSCLC mechanism.',
      'NO_GO_rule':'If author S4 missing/unusable, full fixed-panel primary tests not evaluable, all primary contrasts null, or validation panel coverage is insufficient, retain null/descriptive findings and do not relabel a hypothesis as validated.',
      'limits':['no member tuning','no raw sequencing download','no environment modification','no new wet experiment','no new computation novelty claim; source paper already compares GI/co-essentiality and complex enrichment'],
    }
    if specfile.exists():
        previous=json.loads(specfile.read_text())
        assert previous['fixed_sets']==SETS, 'Frozen sets changed'
    else: save(specfile,spec)
    print('SPEC_FROZEN',sha(specfile),flush=True)
    manifest=[]
    for name in ['File_S1.xlsx','File_S4.xlsx','File_S6.xlsx']:
        url=BASE+'assets/data_files/'+name
        if name=='File_S4.xlsx':
            url='https://media.githubusercontent.com/media/csbio/GIN/main/Data%20Files/File_S4.xlsx'
        elif name=='File_S6.xlsx':
            url='https://raw.githubusercontent.com/csbio/GIN/main/Data%20Files/File_S6.xlsx'
        path=raw/name
        if path.exists():
            row={'url':url,'path':str(path),'bytes':path.stat().st_size,'sha256':sha(path),'existing':True}
            # A failed/interrupted stream must never be reused as a complete input.
            check=openpyxl.load_workbook(path,read_only=True,data_only=True)
            check.close()
        else:
            print('DOWNLOADING',name,flush=True); row=fetch(url,path)
        manifest.append(row); save(out/'download_manifest.json',manifest)
        if name=='File_S4.xlsx':
            assert row['sha256']=='9df16c51264e9bfc8bba9620ccb0eb5cdd308a5d053c1fcebe0152ca07ea57b6', 'Git LFS SHA256 mismatch'
        print('ACQUIRED',name,row['bytes'],row['sha256'],flush=True)
    inventory={}
    for name in ['File_S1.xlsx','File_S4.xlsx','File_S6.xlsx']:
        print('INSPECTING_HEADERS',name,flush=True)
        book=openpyxl.load_workbook(raw/name,read_only=True,data_only=True)
        inventory[name]={}
        for ws in book:
            rows=list(ws.iter_rows(min_row=1,max_row=4,values_only=True))
            inventory[name][ws.title]={'rows':ws.max_row,'cols':ws.max_column,'first_rows':rows}
        book.close()
    save(out/'workbook_inventory.json',inventory)
    save(out/'runtime.json',{'python':platform.python_version(),'platform':platform.platform(),'host':platform.node()})
    print(json.dumps(inventory,default=str),flush=True)

if __name__=='__main__': main()
