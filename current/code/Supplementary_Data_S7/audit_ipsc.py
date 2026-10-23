"""Independent workbook-cell extraction audit using openpyxl and scalar arithmetic."""
from pathlib import Path
import argparse,csv,json,math
from openpyxl import load_workbook
ap=argparse.ArgumentParser();ap.add_argument('--input',type=Path,required=True);ap.add_argument('--results',type=Path,required=True);a=ap.parse_args()
ws=load_workbook(a.input,data_only=True,read_only=True).active
rows=list(ws.iter_rows(values_only=True));gene_col=list(rows[2]).index('Gene Name');lookup={}
for r in rows[3:]:lookup.setdefault(r[gene_col],[]).append(r)
errors=[];n=0
with (a.results/'fixed_panel_KO_results.csv').open() as f:
    for record in csv.DictReader(f):
        assert len(lookup[record['gene']])==1;r=lookup[record['gene']][0]
        ko={'ARID1A':23,'ARID2':25,'PBRM1':27}[record['query']];rna={'ARID1A':32,'ARID2':38,'PBRM1':44}[record['query']]
        for key,j in [('WT_A',21),('WT_B',22),('KO_A',ko),('KO_B',ko+1),('RNA_log2FC',rna),('RNA_source_padj',rna+4)]:
            assert math.isclose(float(record[key]),float(r[j]),rel_tol=1e-12,abs_tol=1e-20),(key,record['gene'])
        fc=math.log2((float(r[ko])+float(r[ko+1]))/(float(r[21])+float(r[22])))
        errors.append(abs(fc-float(record['protein_log2FC'])));n+=1
assert n==39 and max(errors)<1e-12
report={'records_verified':n,'protein_log2FC_max_error':max(errors),'WT_KO_and_source_RNA_cells':'all exact within floating tolerance','method':'Independent openpyxl original cell indices, scalar sums and math.log2; no pairwise protein P values'}
(a.results/'independent_audit.json').write_text(json.dumps(report,indent=2));print(json.dumps(report,indent=2))
